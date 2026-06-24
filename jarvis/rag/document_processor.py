"""
PDF and Document Processing for JARVIS RAG system.
Handles PDF, TXT, DOCX, and Markdown file ingestion.
"""

import re
import hashlib
import logging
from pathlib import Path
from typing import List, Dict, Optional, Any
from dataclasses import dataclass, field
from datetime import datetime

logger = logging.getLogger(__name__)


@dataclass
class DocumentChunk:
    """A chunk of text from a document."""
    id: str
    content: str
    metadata: Dict[str, Any] = field(default_factory=dict)
    source: str = ""
    page: Optional[int] = None
    chunk_index: int = 0


@dataclass
class ProcessedDocument:
    """A processed document with chunks."""
    title: str
    source_path: str
    file_type: str
    chunks: List[DocumentChunk] = field(default_factory=list)
    metadata: Dict[str, Any] = field(default_factory=dict)
    created_at: datetime = field(default_factory=datetime.now)


class DocumentProcessor:
    """
    Process various document formats for RAG ingestion.
    
    Supported formats:
    - PDF (via PyMuPDF or PyPDF2)
    - TXT (plain text)
    - Markdown (md)
    - DOCX (via python-docx)
    """

    DEFAULT_CHUNK_SIZE = 500
    DEFAULT_CHUNK_OVERLAP = 50

    def __init__(self, chunk_size: int = DEFAULT_CHUNK_SIZE, chunk_overlap: int = DEFAULT_CHUNK_OVERLAP):
        self.chunk_size = chunk_size
        self.chunk_overlap = chunk_overlap

    def _generate_chunk_id(self, source: str, chunk_index: int) -> str:
        """Generate unique chunk ID."""
        hash_input = f"{source}_{chunk_index}_{datetime.now().timestamp()}"
        return hashlib.md5(hash_input.encode()).hexdigest()[:16]

    def _split_into_chunks(self, text: str, source: str, metadata: Dict[str, Any]) -> List[DocumentChunk]:
        """Split text into overlapping chunks."""
        if not text.strip():
            return []

        chunks = []
        start = 0
        text_length = len(text)
        chunk_index = 0

        while start < text_length:
            end = min(start + self.chunk_size, text_length)
            
            if end < text_length:
                sentence_break = max(
                    text.rfind('. ', start, end),
                    text.rfind('! ', start, end),
                    text.rfind('? ', start, end),
                    text.rfind('\n', start, end)
                )
                if sentence_break > start + self.chunk_size // 2:
                    end = sentence_break + 2

            chunk_text = text[start:end].strip()
            if chunk_text:
                chunk = DocumentChunk(
                    id=self._generate_chunk_id(source, chunk_index),
                    content=chunk_text,
                    metadata=metadata.copy(),
                    source=source,
                    chunk_index=chunk_index
                )
                chunks.append(chunk)
                chunk_index += 1

            start = end - self.chunk_overlap
            if start <= (chunk_index - 1) * (self.chunk_size - self.chunk_overlap) if chunks else start:
                start = start + self.chunk_size if chunks else start

        return chunks

    async def process_file(self, file_path: Path) -> ProcessedDocument:
        """Process a document file and return chunks."""
        if not file_path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")

        suffix = file_path.suffix.lower()
        
        if suffix == '.pdf':
            return await self._process_pdf(file_path)
        elif suffix == '.txt':
            return await self._process_txt(file_path)
        elif suffix in ['.md', '.markdown']:
            return await self._process_markdown(file_path)
        elif suffix == '.docx':
            return await self._process_docx(file_path)
        else:
            raise ValueError(f"Unsupported file type: {suffix}")

    async def _process_pdf(self, file_path: Path) -> ProcessedDocument:
        """Process a PDF file."""
        metadata = {"file_name": file_path.name, "file_size": file_path.stat().st_size, "file_type": "pdf"}
        chunks = []
        page_num = 0
        
        try:
            import fitz
            doc = fitz.open(str(file_path))
            metadata["page_count"] = len(doc)
            
            for page in doc:
                page_num += 1
                text = page.get_text()
                if text.strip():
                    page_chunks = self._split_into_chunks(text, str(file_path), {**metadata, "page": page_num})
                    for chunk in page_chunks:
                        chunk.page = page_num
                    chunks.extend(page_chunks)
            doc.close()
            
        except ImportError:
            try:
                import PyPDF2
                with open(file_path, 'rb') as f:
                    reader = PyPDF2.PdfReader(f)
                    metadata["page_count"] = len(reader.pages)
                    for page in reader.pages:
                        page_num += 1
                        text = page.extract_text()
                        if text and text.strip():
                            page_chunks = self._split_into_chunks(text, str(file_path), {**metadata, "page": page_num})
                            for chunk in page_chunks:
                                chunk.page = page_num
                            chunks.extend(page_chunks)
            except ImportError:
                raise ImportError("Install PyMuPDF or PyPDF2 for PDF support")

        return ProcessedDocument(title=file_path.stem, source_path=str(file_path), file_type="pdf", chunks=chunks, metadata=metadata)

    async def _process_txt(self, file_path: Path) -> ProcessedDocument:
        """Process a plain text file."""
        metadata = {"file_name": file_path.name, "file_size": file_path.stat().st_size, "file_type": "txt"}
        text = file_path.read_text(encoding='utf-8', errors='ignore')
        chunks = self._split_into_chunks(text, str(file_path), metadata)
        return ProcessedDocument(title=file_path.stem, source_path=str(file_path), file_type="txt", chunks=chunks, metadata=metadata)

    async def _process_markdown(self, file_path: Path) -> ProcessedDocument:
        """Process a Markdown file."""
        metadata = {"file_name": file_path.name, "file_size": file_path.stat().st_size, "file_type": "markdown"}
        text = file_path.read_text(encoding='utf-8', errors='ignore')
        text = self._clean_markdown(text)
        chunks = self._split_into_chunks(text, str(file_path), metadata)
        return ProcessedDocument(title=file_path.stem, source_path=str(file_path), file_type="md", chunks=chunks, metadata=metadata)

    def _clean_markdown(self, text: str) -> str:
        """Remove markdown formatting from text."""
        text = re.sub(r'```[\s\S]*?```', '', text)
        text = re.sub(r'`[^`]+`', '', text)
        text = re.sub(r'\[([^\]]+)\]\([^\)]+\)', r'\1', text)
        text = re.sub(r'!\[([^\]]*)\]\([^\)]+\)', '', text)
        text = re.sub(r'^#+\s+', '', text, flags=re.MULTILINE)
        text = re.sub(r'\*\*?([^\*]+)\*\*?', r'\1', text)
        text = re.sub(r'^[-*_]{3,}$', '', text, flags=re.MULTILINE)
        return text

    async def _process_docx(self, file_path: Path) -> ProcessedDocument:
        """Process a DOCX file."""
        try:
            from docx import Document
        except ImportError:
            raise ImportError("Install python-docx for DOCX support")
        
        metadata = {"file_name": file_path.name, "file_size": file_path.stat().st_size, "file_type": "docx"}
        doc = Document(file_path)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        text = "\n\n".join(paragraphs)
        chunks = self._split_into_chunks(text, str(file_path), metadata)
        return ProcessedDocument(title=file_path.stem, source_path=str(file_path), file_type="docx", chunks=chunks, metadata=metadata)


class StudyAssistant:
    """Study assistant for generating summaries, questions, and exam prep."""

    def __init__(self, document_processor: DocumentProcessor):
        self.processor = document_processor

    async def summarize_document(self, file_path: Path, max_length: int = 500) -> str:
        """Generate a summary of the document."""
        doc = await self.processor.process_file(file_path)
        
        summary_parts = []
        current_length = 0
        
        for chunk in doc.chunks[:5]:
            if current_length + len(chunk.content) > max_length:
                break
            summary_parts.append(chunk.content)
            current_length += len(chunk.content)
        
        summary = "\n\n".join(summary_parts)
        
        if len(summary) > max_length:
            summary = summary[:max_length].rsplit(' ', 1)[0] + "..."
        
        return summary

    async def generate_questions(self, file_path: Path, num_questions: int = 5, question_type: str = "general") -> List[str]:
        """Generate practice questions from document."""
        doc = await self.processor.process_file(file_path)
        content = "\n\n".join([c.content for c in doc.chunks[:10]])
        questions = []
        key_terms = self._extract_key_terms(content)
        
        for term in key_terms[:num_questions]:
            if len(questions) >= num_questions:
                break
            questions.append(f"What is '{term}'? (Context: found in document)")
        
        general_questions = [
            "What are the main topics covered in this document?",
            "Summarize the key points of this document.",
            "What new information did you learn from this document?",
        ]
        
        while len(questions) < num_questions and general_questions:
            questions.append(general_questions.pop(0))
        
        return questions[:num_questions]

    def _extract_key_terms(self, text: str, max_terms: int = 10) -> List[str]:
        """Extract key terms from text."""
        stop_words = {'the', 'a', 'an', 'and', 'or', 'but', 'in', 'on', 'at', 'to', 'for',
                      'of', 'with', 'by', 'from', 'is', 'was', 'are', 'be', 'been', 'have',
                      'has', 'had', 'do', 'does', 'did', 'this', 'that', 'these', 'those'}
        
        words = re.findall(r'\b[a-zA-Z]{5,}\b', text.lower())
        words = [w for w in words if w not in stop_words]
        word_freq = {}
        for word in words:
            word_freq[word] = word_freq.get(word, 0) + 1
        
        sorted_words = sorted(word_freq.items(), key=lambda x: x[1], reverse=True)
        return [word for word, _ in sorted_words[:max_terms]]

    async def generate_exam_notes(self, file_path: Path, format_type: str = "bullet") -> str:
        """Generate exam revision notes from document."""
        doc = await self.processor.process_file(file_path)
        notes = [f"# Revision Notes: {doc.title}\n"]
        
        for i, chunk in enumerate(doc.chunks[:20]):
            content = chunk.content.strip()
            
            if format_type == "bullet":
                lines = content.split('\n')
                for line in lines[:3]:
                    if line.strip():
                        notes.append(f"• {line.strip()}")
            elif format_type == "outline":
                notes.append(f"{i+1}. {content[:200]}...")
            else:
                notes.append(f"---\nQ: {content[:150]}...\nA: (See full document)")
        
        return "\n".join(notes)
